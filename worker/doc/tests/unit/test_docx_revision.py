import json
from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.errors import WorkerError
from app.main import app
from app.services.docx_revision import TABLE_TITLE, DocumentEdit, revise

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def source() -> bytes:
    document = Document()
    document.add_heading("취업규칙", level=1)
    rule = document.add_paragraph(style="List Number")
    rule.add_run("연장근로는 ")
    rule.add_run("1주 12시간").bold = True
    rule.add_run("을 초과할 수 없다.")
    document.add_paragraph("육아휴직은 나누어 사용할 수 없다.")
    document.add_paragraph("부칙: 이 규칙은 2026년 1월 1일부터 시행한다.")
    table = document.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = "경조휴가는 본인 결혼 5일이다."
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def paragraphs(content: bytes) -> list:
    return Document(BytesIO(content)).paragraphs


def test_replaces_only_the_matched_text_and_keeps_the_paragraph_style() -> None:
    revised = revise(source(), [DocumentEdit(before="육아휴직은 나누어 사용할 수 없다.", after="육아휴직은 3회까지 나누어 사용할 수 있다.")])

    texts = [paragraph.text for paragraph in paragraphs(revised)]
    assert "육아휴직은 3회까지 나누어 사용할 수 있다." in texts
    assert "육아휴직은 나누어 사용할 수 없다." not in texts[:4]
    assert paragraphs(revised)[0].style.name == "Heading 1"
    assert paragraphs(revised)[1].style.name == "List Number"


def test_a_match_across_runs_takes_the_first_runs_format_and_leaves_the_rest() -> None:
    revised = revise(source(), [DocumentEdit(before="1주 12시간을 초과할", after="1주 12시간(월 52시간 한도)을 초과할")])

    rule = paragraphs(revised)[1]
    assert rule.text == "연장근로는 1주 12시간(월 52시간 한도)을 초과할 수 없다."
    assert rule.style.name == "List Number"
    assert rule.runs[0].text == "연장근로는 "


def test_whitespace_is_compared_loosely_and_table_cells_are_searched() -> None:
    revised = revise(source(), [DocumentEdit(before="경조휴가는\n본인   결혼 5일이다.", after="경조휴가는 본인 결혼 7일이다.")])

    assert Document(BytesIO(revised)).tables[0].rows[0].cells[0].text == "경조휴가는 본인 결혼 7일이다."


def test_appends_a_comparison_table_of_the_changed_paragraphs() -> None:
    revised = revise(source(), [
        DocumentEdit(before="나누어 사용할 수 없다.", after="3회까지 나누어 사용할 수 있다."),
        DocumentEdit(before="2026년 1월 1일", after="2026년 7월 1일"),
    ])

    document = Document(BytesIO(revised))
    assert TABLE_TITLE in [paragraph.text for paragraph in document.paragraphs]
    rows = [[cell.text for cell in row.cells] for row in document.tables[-1].rows]
    assert rows == [
        ["현행", "개정"],
        ["육아휴직은 나누어 사용할 수 없다.", "육아휴직은 3회까지 나누어 사용할 수 있다."],
        ["부칙: 이 규칙은 2026년 1월 1일부터 시행한다.", "부칙: 이 규칙은 2026년 7월 1일부터 시행한다."],
    ]


def test_two_edits_in_one_paragraph_keep_the_first_and_last_text() -> None:
    revised = revise(source(), [
        DocumentEdit(before="2026년", after="2027년"),
        DocumentEdit(before="1월 1일", after="3월 1일"),
    ])

    rows = [[cell.text for cell in row.cells] for row in Document(BytesIO(revised)).tables[-1].rows]
    assert rows[1:] == [["부칙: 이 규칙은 2026년 1월 1일부터 시행한다.", "부칙: 이 규칙은 2027년 3월 1일부터 시행한다."]]


@pytest.mark.parametrize(
    ("before", "code"),
    [("없는 문장이다.", "EDIT_NOT_FOUND"), ("수 없다.", "EDIT_AMBIGUOUS"), ("시행한다. 경조휴가는", "EDIT_NOT_FOUND")],
)
def test_refuses_text_that_is_missing_repeated_or_spans_paragraphs(before: str, code: str) -> None:
    with pytest.raises(WorkerError) as error:
        revise(source(), [DocumentEdit(before=before, after="x")])

    assert error.value.code == code


def test_refuses_a_file_that_is_not_docx() -> None:
    with pytest.raises(WorkerError) as error:
        revise(b"plain text", [DocumentEdit(before="a", after="b")])

    assert error.value.status_code == 422


def test_revisions_endpoint_returns_the_revised_docx(monkeypatch) -> None:
    monkeypatch.setenv("INTERNAL_API_KEY", "test-key")
    headers = {"X-Internal-Api-Key": "test-key", "X-Tenant-Id": "tenant-a"}
    edits = json.dumps([{"before": "1월 1일", "after": "7월 1일"}], ensure_ascii=False)

    response = TestClient(app).post("/api/v1/documents/revisions", headers=headers,
                                    files={"file": ("rules.docx", source(), DOCX)}, data={"edits": edits})

    assert response.status_code == 200
    assert response.headers["content-type"] == DOCX
    assert "부칙: 이 규칙은 2026년 7월 1일부터 시행한다." in [paragraph.text for paragraph in paragraphs(response.content)]


def test_revisions_endpoint_reports_a_missing_edit_and_bad_input(monkeypatch) -> None:
    monkeypatch.setenv("INTERNAL_API_KEY", "test-key")
    headers = {"X-Internal-Api-Key": "test-key", "X-Tenant-Id": "tenant-a"}
    client = TestClient(app)

    missing = client.post("/api/v1/documents/revisions", headers=headers, files={"file": ("rules.docx", source(), DOCX)},
                          data={"edits": json.dumps([{"before": "없는 문장", "after": "x"}], ensure_ascii=False)})
    malformed = client.post("/api/v1/documents/revisions", headers=headers, files={"file": ("rules.docx", source(), DOCX)},
                            data={"edits": "[{\"from\": \"a\"}]"})
    unauthorized = client.post("/api/v1/documents/revisions", headers={**headers, "X-Internal-Api-Key": "wrong"},
                               files={"file": ("rules.docx", source(), DOCX)}, data={"edits": "[]"})

    assert (missing.status_code, missing.json()["code"]) == (422, "EDIT_NOT_FOUND")
    assert (malformed.status_code, malformed.json()["code"]) == (400, "VALIDATION_ERROR")
    assert unauthorized.status_code == 401
