"""원본 docx의 서식을 유지한 채 지정한 문장만 바꾸고, 끝에 신구조문 대비표를 붙인다.

문단 스타일·번호·표·머리글은 건드리지 않고, 바꿀 문장이 걸친 run의 글자만 고친다(바꾼 글은 시작 run의 서식을 따른다).
바꿀 문장은 문서 전체에서 정확히 한 곳(한 문단 안)에 있어야 한다 — 없거나 여러 곳이면 엉뚱한 곳을 고치지 않게 실패한다.
공백(줄바꿈·연속 공백)은 비교할 때만 하나로 본다. 검색 조각처럼 공백이 정리된 글을 그대로 넘겨도 찾게 하기 위해서다.
"""

from collections.abc import Iterator
from io import BytesIO

from docx import Document
from docx.document import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from pydantic import BaseModel, ConfigDict, Field

from app.errors import WorkerError

MAX_EDITS = 50
MAX_EDIT_TEXT = 2_000
TABLE_TITLE = "신구조문 대비표"


class DocumentEdit(BaseModel):
    model_config = ConfigDict(extra="forbid")

    before: str = Field(min_length=1, max_length=MAX_EDIT_TEXT)
    after: str = Field(max_length=MAX_EDIT_TEXT)


def revise(content: bytes, edits: list[DocumentEdit]) -> bytes:
    if not 1 <= len(edits) <= MAX_EDITS:
        raise WorkerError(400, "VALIDATION_ERROR", f"edits must have 1 to {MAX_EDITS} items")
    try:
        document = Document(BytesIO(content))
    except Exception as error:  # python-docx는 깨진 zip·xml에 여러 예외를 낸다.
        raise WorkerError(422, "UNPROCESSABLE_DOCUMENT", "Source is not a valid docx document") from error

    paragraphs = list(_paragraphs(document))
    # 문단별 (현행, 개정). 한 문단을 여러 번 고치면 처음 글과 마지막 글을 남긴다.
    changes: dict[int, tuple[str, str]] = {}
    for number, edit in enumerate(edits, start=1):
        index, original = _replace_once(paragraphs, edit, number)
        changes[index] = (changes[index][0] if index in changes else original, _text(paragraphs[index]))

    _append_comparison_table(document, [changes[index] for index in sorted(changes)])
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _replace_once(paragraphs: list[Paragraph], edit: DocumentEdit, number: int) -> tuple[int, str]:
    """before가 정확히 한 곳에 있으면 바꾸고, 그 문단의 위치와 바꾸기 전 글을 돌려준다."""
    target, _ = _normalize(edit.before)
    found: list[tuple[int, int]] = []
    for index, paragraph in enumerate(paragraphs):
        normalized, _ = _normalize(_text(paragraph))
        start = normalized.find(target)
        while start != -1:
            found.append((index, start))
            start = normalized.find(target, start + 1)
    if not found:
        raise WorkerError(422, "EDIT_NOT_FOUND", f"edit {number}: before text was not found in a single paragraph")
    if len(found) > 1:
        raise WorkerError(422, "EDIT_AMBIGUOUS", f"edit {number}: before text appears {len(found)} times")

    index, start = found[0]
    paragraph = paragraphs[index]
    original = _text(paragraph)
    _, positions = _normalize(original)
    _replace_range(paragraph, positions[start], positions[start + len(target) - 1] + 1, edit.after)
    return index, original


def _replace_range(paragraph: Paragraph, start: int, end: int, replacement: str) -> None:
    """run 글자를 이어 붙인 위치 [start, end)를 replacement로 바꾼다. 시작 run에 넣고, 걸친 나머지 run에서는 지운다."""
    offset = 0
    first = None
    for run in paragraph.runs:
        text = run.text
        run_start, run_end = offset, offset + len(text)
        offset = run_end
        if run_end <= start or run_start >= end:
            continue
        keep_before = text[: max(start - run_start, 0)] if first is None else ""
        keep_after = text[end - run_start:] if end <= run_end else ""
        if first is None:
            first = run
            run.text = keep_before + replacement + keep_after
        else:
            run.text = keep_after


def _normalize(text: str) -> tuple[str, list[int]]:
    """공백 묶음을 공백 하나로 줄이고 앞뒤 공백을 뺀 글과, 그 글자마다 원래 글에서의 위치를 돌려준다."""
    characters: list[str] = []
    positions: list[int] = []
    for position, character in enumerate(text):
        if character.isspace():
            if characters and characters[-1] != " ":
                characters.append(" ")
                positions.append(position)
        else:
            characters.append(character)
            positions.append(position)
    if characters and characters[-1] == " ":
        characters.pop()
        positions.pop()
    return "".join(characters), positions


def _text(paragraph: Paragraph) -> str:
    # paragraph.text는 하이퍼링크 글자도 포함하지만 runs에는 없다. 고칠 수 있는 글자(runs)만 비교한다.
    return "".join(run.text for run in paragraph.runs)


def _paragraphs(document: DocxDocument) -> Iterator[Paragraph]:
    yield from document.paragraphs
    for table in document.tables:
        yield from _table_paragraphs(table)


def _table_paragraphs(table: Table) -> Iterator[Paragraph]:
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from _table_paragraphs(nested)


def _append_comparison_table(document: DocxDocument, changes: list[tuple[str, str]]) -> None:
    document.add_page_break()
    try:
        document.add_heading(TABLE_TITLE, level=1)
    except KeyError:  # 원본에 제목 스타일이 없으면 굵은 문단으로 쓴다.
        document.add_paragraph().add_run(TABLE_TITLE).bold = True
    table = document.add_table(rows=1, cols=2)
    try:
        table.style = "Table Grid"
    except (KeyError, ValueError):
        pass
    table.rows[0].cells[0].text = "현행"
    table.rows[0].cells[1].text = "개정"
    for before, after in changes:
        cells = table.add_row().cells
        cells[0].text = before
        cells[1].text = after
